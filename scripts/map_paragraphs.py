import docx

doc = docx.Document(r'S:\Laporan Magang\LAPORAN MAGANG SHELLA DRAFT 2.docx')
with open('paragraphs_map.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        txt = p.text[:80].replace('\n', ' ')
        f.write(f"{i:3d} | [{p.style.name:12s}] | {txt}\n")

print(f"Mapped {len(doc.paragraphs)} paragraphs.")
