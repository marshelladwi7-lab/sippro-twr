# -*- coding: utf-8 -*-
import os
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

TARGET = r"S:\Laporan Magang\LAPORAN MAGANG SHELLA DRAFT 3_PERFECTED.docx"

def main():
    if not os.path.exists(TARGET):
        print(f"[ERROR] Target file not found: {TARGET}")
        return

    doc = docx.Document(TARGET)
    print(f"=== VERIFICATION OF: {os.path.basename(TARGET)} ===")
    print(f"File Size: {os.path.getsize(TARGET):,} bytes")
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")

    # 1. Total Word Count
    total_words = sum(len(p.text.split()) for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                total_words += sum(len(p.text.split()) for p in cell.paragraphs)
    print(f"Total Word Count: {total_words:,} words")

    # 2. Embedded Drawings / Images
    p_drawings = []
    for i, p in enumerate(doc.paragraphs):
        drawings = p._p.xpath('.//w:drawing')
        if drawings:
            p_drawings.append((i, p.text[:50]))
    print(f"\nEmbedded Drawings / Images Found: {len(p_drawings)}")

    # 3. Figures Audit
    print("\nFigures Found:")
    figures = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t.startswith("Gambar "):
            figures.append((i, t))
            print(f"  P{i:3d}: {t[:75]}")
    print(f"Total Figures: {len(figures)}")

    # 4. Tables Audit
    print("\nTables Summary:")
    for i, t in enumerate(doc.tables):
        header = [c.text.strip().replace('\n', ' ')[:30] for c in t.rows[0].cells]
        print(f"  Table {i+1}: {len(t.rows)} rows x {len(t.columns)} cols -> {header[:3]}")

    # 5. Markdown asterisks
    double_stars = sum(1 for p in doc.paragraphs if '**' in p.text)
    print(f"\nParagraphs with '**': {double_stars}")

    # 6. Awkward & Forbidden Words
    print("\nForbidden / Awkward Words Check:")
    words_to_check = [
        'bilah', 'penapisan', 'ingesti', 'sanitasi',
        'haircut', 'matriks penyesuaian', 'mesin likuidasi', 'executive bank summary'
    ]
    for w in words_to_check:
        matches = [i for i, p in enumerate(doc.paragraphs) if w in p.text.lower()]
        print(f"  '{w}': {len(matches)} matches {matches if matches else ''}")

    # 7. Alignment Distribution
    print("\nBody Alignment Distribution (from BAB I onwards):")
    aligns = {}
    bab1_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith("BAB I"))
    for p in doc.paragraphs[bab1_idx:]:
        align_str = str(p.alignment)
        aligns[align_str] = aligns.get(align_str, 0) + 1
    for k, v in aligns.items():
        print(f"  {k}: {v}")

    # 8. Margins
    print("\nSection Margins:")
    for i, s in enumerate(doc.sections):
        print(f"  Section {i+1}: Top={s.top_margin.cm:.2f}cm, Bottom={s.bottom_margin.cm:.2f}cm, Left={s.left_margin.cm:.2f}cm, Right={s.right_margin.cm:.2f}cm")

if __name__ == "__main__":
    main()
