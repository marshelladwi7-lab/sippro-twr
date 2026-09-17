# -*- coding: utf-8 -*-
import docx

doc_path = r"S:\Laporan Magang\DOKUMENTASI_SIMULASI_DAN_DEMO_SIPPRO_TWR_FIXED.docx"
doc = docx.Document(doc_path)

print("=== VERIFICATION REPORT ===")
print("File:", doc_path)
print("Total Paragraphs:", len(doc.paragraphs))
print("Total Tables:", len(doc.tables))

# 1. Check for **
star_count = 0
for p in doc.paragraphs:
    if "**" in p.text:
        star_count += 1
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if "**" in p.text:
                    star_count += 1
print("1. Double asterisks (**) count:", star_count, "-> PASS" if star_count == 0 else "-> FAIL")

# 2. Check for forbidden non-live terms
forbidden = [
    "executive bank summary",
    "matriks penyesuaian",
    "mesin likuidasi",
    "haircut",
    "bilah",
    "ingesti",
    "sanitasi",
    "penapisan"
]

forbidden_hits = {}
all_clear = True
for term in forbidden:
    hits = 0
    for p in doc.paragraphs:
        if term.lower() in p.text.lower():
            hits += 1
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if term.lower() in p.text.lower():
                        hits += 1
    forbidden_hits[term] = hits
    if hits > 0:
        all_clear = False

print("2. Forbidden / Awkward words check:")
for term, hits in forbidden_hits.items():
    print(f"   - '{term}': {hits} hits")
print("   -> Overall Non-live / Awkward terms check:", "PASS" if all_clear else "FAIL")

# 3. Check embedded images
images_count = 0
for rel in doc.part.rels.values():
    if "image" in rel.target_ref:
        images_count += 1
print("3. Total embedded images:", images_count)

# 4. List Headings
print("\n4. Document Headings:")
for p in doc.paragraphs:
    if p.style.name.startswith("Heading") or p.text.startswith(("BAGIAN", "RINGKASAN", "3.", "4.", "5.")):
        if len(p.text) < 100:
            print("   *", p.text)
